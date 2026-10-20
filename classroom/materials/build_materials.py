from pathlib import Path
import json, html, shutil
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parent
ref=json.loads((root/'game-reference.json').read_text())
shutil.copyfile(root.parent/'artifacts/student-city.png',root/'student-city.png')
shutil.copyfile(root.parent/'artifacts/teacher-final.png',root/'teacher-chart.png')
def table(head,rows):
 return '<table><thead><tr>'+''.join('<th>'+x+'</th>' for x in head)+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+str(x)+'</td>' for x in row)+'</tr>' for row in rows)+'</tbody></table>'
def box(title,text):return f'<aside><b>{title}</b><p>{text}</p></aside>'
def lines(n=2):return '<div class="lines">'+''.join('<p>　</p>' for _ in range(n))+'</div>'
def section(k,t,body):return {'k':k,'title':t,'body':body}
student=[]
student.append(section('사전 학습 01','인구가 줄어도, 살기 좋은 도시일까?', '''<p class="lead">여러분은 다온시의 운영자입니다. 같은 도시에서 출발하지만, <b>어디에 무엇을 먼저 투자하는지</b>에 따라 서로 다른 미래를 만듭니다.</p>
<div class="goals"><b>오늘의 학습 목표</b><p>① 인구 문제의 원인과 지역별 차이를 설명한다.<br>② 정책의 효과와 한계를 주거·일자리·서비스·재정의 관계로 설명한다.<br>③ 인구 증가와 생활 여건 개선을 구분해 지속가능한 도시를 평가한다.</p></div>
<h2>인구 변화는 두 흐름의 합</h2><div class="formula">인구 증감 = 자연적 증감 + 사회적 증감<br><small>자연적 증감: 출생 − 사망　｜　사회적 증감: 전입 − 전출</small></div><p>출생이 사망보다 적어도 전입이 많으면 인구가 늘 수 있습니다. 반대로 출생이 있어도 더 많은 사람이 다른 지역으로 떠나면 인구가 줄 수 있습니다.</p>
'''+table(['개념','뜻과 살펴볼 점'],[
['저출생','출생 규모가 낮은 현상. 합계출산율은 현재 연령별 출산 수준이 유지된다고 가정할 때 여성 1명이 평생 낳을 것으로 예상되는 평균 출생아 수이며, 출생아 수 자체와 다릅니다.'],
['고령화','전체 인구에서 고령 인구가 차지하는 비중이 높아지는 현상. 기대수명 증가와 낮은 출생 수준 등이 함께 영향을 줍니다.'],
['인구 이동·지역 격차','일자리·주거·교육·서비스 등을 고려한 이동이 지역별 인구와 생활 여건의 차이와 연결됩니다.'],
['인구 과잉','빠른 인구 증가에 식량·일자리·기반 시설·주택 공급이 충분히 대응하지 못하는 문제를 살펴봅니다. 사람 수만으로 생활 여건을 판단하지 않습니다.']])+
'<h2>선택하기 전에 던질 질문</h2><p>이 도시의 문제는 <b>사람의 수</b>일까요, <b>일자리·주택·서비스의 배치와 이용 조건</b>일까요? 두 문제가 어떻게 연결되는지 찾아봅시다.</p>'+box('기억할 문장','시설을 지었다는 사실과 주민이 실제로 이용할 수 있다는 사실은 다릅니다.')))
student.append(section('사전 학습 02','문제와 해결을 함께 읽기', '<p>제공된 사진의 인구 문제와 해결 방안을 수업용으로 재구성했습니다. 국가·지역의 상황은 다양하므로 아래 설명을 모든 지역에 똑같이 적용하지 않습니다.</p>'+table(['문제','원인과 영향','대응 방향'],[
['빠른 인구 증가·도시 집중','공급 속도보다 수요가 빠르게 늘면 식량·일자리·주택·시설 부족이 나타날 수 있습니다.','인구 부양력(자원·기술·생산 조건이 지탱할 수 있는 인구 규모) 개선, 기반 시설 확충, 중소 도시·농촌의 생활 여건 개선.'],
['저출생','양육·주거·교육 부담, 일과 돌봄의 병행 어려움, 고용 여건과 가족·삶에 대한 가치관 변화 등이 복합적으로 작용합니다.','양육 지원, 보육 확충, 출산·육아휴직과 일·생활 균형, 성평등한 돌봄 분담 등.'],
['고령화','노동 공급과 돌봄·의료 수요, 사회 보장 지출의 구성이 달라집니다.','의료·이동 접근성, 연금 등 노후 소득 보장, 정년·고용 제도와 직업 훈련, 지역 돌봄과 시설 재편.'],
['지역 인구 격차','일자리·교육·서비스의 지역 차이와 이동이 맞물립니다. 집중 지역에는 주거비·교통 부담, 유출 지역에는 시설 유지의 어려움이 생길 수 있습니다.','지역 일자리와 정주 여건, 교육·의료·이동 서비스, 인구 규모에 맞춘 시설 배치 개선.']])+
'<h2>사진 속 표현을 읽을 때 주의할 점</h2><ul><li>사진은 인구 과잉을 일부 개발도상 지역, 저출생·고령화를 일부 선진 지역의 경향과 연결합니다. <b>국가별 차이가 있으므로 고정된 구분은 아닙니다.</b></li><li>사진의 “2018년 고령 사회 진입”은 과거 시점의 서술입니다. 오늘의 단계·비율은 별도의 최신 통계 확인이 필요합니다.</li><li>저출생을 여성의 사회 참여 탓으로 단정하지 않습니다. 일과 돌봄을 함께 할 수 있는 제도와 부담의 분담을 살펴봅니다.</li><li>인구 증가 억제 정책이 소개되더라도 개인의 재생산에 관한 권리와 선택을 함께 고려해야 합니다.</li></ul>'+box('정책을 판단하는 기준','세대 간 정의: 현재와 미래 세대의 부담·혜택을 함께 고려하기.<br>양성평등·일생활 균형: 돌봄 책임을 특정 성별에만 맡기지 않기.<br>노인의 사회 참여: 고령 주민을 부담으로만 보지 않고 경험·권리·참여를 존중하기.')))
student.append(section('게임 안내 03','접속부터 첫 결과까지', '''<div class="name">학년·반 __________　조 이름 __________　운영자 __________　분석가 __________</div>
<ol><li><b>교사가 알려 준 접속 주소</b>를 엽니다. 기존 개인용 게임 주소와 다릅니다.</li><li>참가 코드와 조 이름을 입력합니다. 실명 대신 “1조”처럼 씁니다.</li><li>교사가 수업을 시작하면 목표와 지역의 문제를 확인합니다.</li><li><b>생활권 선택 → 정책 검토 → 시행 → 다음 해 → 결산</b> 순서로 운영합니다.</li><li>사건이 나오면 대응을 선택합니다. 8턴 뒤 결과를 비교합니다.</li></ol>
<img class="game" src="student-city.png" alt="조 이름과 서버 기록 아래에 도시 지도, 생활권 버튼, 지표, 정책 영역이 배치된 실제 수업용 게임 화면">
'''+table(['어디를 볼까?','어떻게 읽을까?'],[
['상단 핵심 정보','보유 예산은 현재 가진 돈, 연간 수지는 1년의 수입−지출입니다.'],
['생활권과 시설 정보','주택 공급≠실제 입주, 일자리 수≠실제 취업, 시설 정원≠실제 이용입니다.'],
['정책 검토 창','투자비·유지비·완공 시점·직접 효과·효과의 조건을 확인한 뒤 시행합니다.'],
['결산과 지도','공사→완공→운영 변화와 지표가 함께 달라지는지 확인합니다.']])+
'<p class="note">한 조에 조작 기기 하나를 사용합니다. 운영자는 조작과 제안, 분석가는 비용·조건·결과 확인을 맡고 4턴 뒤 교대합니다. 접속이 끊기면 서버 기록으로 복원될 때까지 기다립니다. 수업 중 “참가 정보 지우기”는 누르지 마세요.</p>'))
policy_rows=[]
for p in ref['policies']:
 direct={'houses':'주거 가능 공간 증가. 입주는 일자리·정착 조건에 따라 달라짐.','jobs':'일자리와 돌봄 인력 여건 개선. 노동 공급이 한도.','care':'돌봄 정원 증가·시설 전환. 인력·접근성·수요가 이용을 제한.','health':'의료·돌봄 접근과 이웃 생활권 연결 개선.','integration':'정착 연결 개선. 빈 주택과 일자리 없이는 유입 제한.','consolidate':'기존 시설 재편, 접근·인력 여건 개선과 비용 절감. 정원 축소에 유의.'}[p['id']]
 policy_rows.append([p['name']+'<br><small>'+('선택 생활권' if p['scope']=='district' else '도시 전체')+'</small>',str(p['cost'])+'억',('즉시' if not p['delay'] else str(p['delay'])+'년'),('2억 절감' if p['upkeep']<0 else str(p['upkeep'])+'억'),direct])
student.append(section('게임 안내 04','정책은 효과와 조건을 함께 비교', '<p class="note">아래 비용·기간은 <b>게임 규칙</b>입니다. 실제 정책 비용이나 효과를 뜻하지 않습니다. 1턴은 1년, 한 판은 8년입니다.</p>'+table(['정책 / 범위','투자','준비','연간 유지','직접 변화와 조건'],policy_rows)+
'<h2>운영 규칙</h2><ul><li>한 해에 신규 사업은 1개. 건설 사업은 동시에 1개만 진행합니다.</li><li>투자비는 시행 때 지출합니다. 운영 가능한 사업의 유지비는 이후 연간 결산에 반영됩니다.</li><li>같은 정책·같은 대상에는 중복 투자할 수 없습니다. 일부 운영 정책은 중지·재개할 수 있지만 투자비는 돌려받지 못합니다.</li><li>게임 종료 전에 완공할 수 없는 사업은 시행할 수 없습니다. 뒤늦은 장기 투자는 결과를 관찰할 시간이 짧습니다.</li></ul>'+table(['교사가 선택한 목표','목표를 확인할 때'],[
['정착 기반','누적 순이동 +80명, 취업 +40명 이상. 인구 총수 증가와 같은 뜻은 아닙니다.'],
['생활 여건','돌봄 충족 85%, 의료 접근 78% 이상. 인구 증가 자체가 필수 조건은 아닙니다.'],
['공통 기본 조건','예비금 3억 이상, 확정 사업의 유지비를 반영한 연간 수지 0 이상, 운영비 전액 지급, 기본 돌봄 충족 40% 이상.']])+
box('결과를 해석하는 방법','고정된 출생·이동 가정과 단순한 접근성 규칙을 사용합니다. 의료 접근성은 모형의 지표이며 실제 조사 비율이 아닙니다. 시민·차량은 보조 연출입니다. 게임 목표 달성 여부만으로 수업 성취를 평가하지 않습니다.')))
student.append(section('활동지 05','우리 조의 선택과 중간 점검', '<div class="name">조 이름 __________　시나리오 __________　목표 __________　시드 __________</div><h2>시작: 해결할 문제 하나를 정하자</h2>'+table(['선택한 생활권','문제를 보여주는 지표·지도 상태','먼저 해결할 이유'],[['　<br>　','　<br>　','　<br>　']])+ '<p>우리의 첫 가설: __________ 정책을 __________ 생활권에 시행하면, __________ 조건이 개선되어 __________ 지표가 달라질 것이다.</p>'+lines(1)+
'<h2>중요한 선택 두 번만 기록하자</h2>'+table(['확인할 것','첫 번째 선택','두 번째 선택'],[
['선택한 턴 / 정책 / 대상','　<br>　','　<br>　'],['정책 전 문제를 보여주는 수치','　<br>　','　<br>　'],['투자비 / 유지비 / 완료 시점','　<br>　','　<br>　'],['예상 효과와 필요한 조건','　<br>　','　<br>　'],['실제 결과: 언제, 무엇이 변했나?','　<br>　','　<br>　'],['미룬 대안과 이유','　<br>　','　<br>　']])+
'<h2>4턴: 역할을 바꾸며 점검</h2><p>예상과 다른 점은? <b>공급·인력·접근성·수요·주거·일자리·재정</b> 중 어떤 조건이 영향을 주었는가?</p>'+lines(2)))
student.append(section('활동지 06','다른 도시와 비교하고 현실로 돌아가기', '<div class="name">우리 조 __________　비교 조 __________　비교한 턴 __________</div><p>교사 차트에서 <b>같은 지표·같은 턴</b>을 비교합니다. 초기 시나리오·목표·시드가 같은지도 확인하세요.</p>'+table(['지표','우리 조 초기 → 비교 턴','다른 조 같은 턴'],[
['총인구 / 인구 증감','　','　'],['실제 취업','　','　'],['돌봄 충족 / 의료 접근','　','　'],['보유 예산 / 연간 수지','　','　']])+
'<h2>주장 → 근거 → 이유로 설명하기</h2><p><b>주장:</b> 우리 도시는 __________________________ 면에서 개선되었다.</p><p><b>근거:</b> __________턴의 __________ 지표가 __________에서 __________로 변했다.</p><p><b>이유:</b> __________ 선택이 __________ 조건을 바꾸었기 때문이다.</p>'+lines(1)+
'<p><b>감수한 점:</b> 개선 과정에서 미루거나 해결하지 못한 것은 무엇인가?</p>'+lines(2)+
'<p><b>공통 질문:</b> 인구가 감소했더라도 생활 여건이 좋아졌다고 판단할 수 있는 근거는?</p>'+lines(2)+
'<p><b>현실로 확장:</b> 게임에 충분히 나타나지 않은 실제 인구 문제의 원인·해결 조건 한 가지는?</p>'+lines(2)+
'<div class="exit"><b>나의 한 문장 정리</b><p>지속가능한 인구 정책은 __________________________만 늘리는 것이 아니라, __________________________도 함께 고려해야 한다.</p></div>'))
teacher=[]
teacher.append(section('교사용 01','20명 공개수업 운영안', '<p class="lead">핵심 질문: <b>“같은 도시인데 왜 다른 미래가 만들어질까?”</b></p><p>20명은 2인 1조·10개 도시로 운영합니다. 한 조에 조작 기기 하나, 교사는 집계 화면을 띄웁니다. 한 판 시간은 <b>첫 플레이 20~30분 예상</b>이며 학생 실측값이 아닙니다.</p>'+table(['50분 기준','활동','관찰·발문'],[
['0~5분','사전 학습 핵심 확인','인구 감소와 삶의 질 악화는 항상 같은가?'],['5~9분','교사 짧은 시연·조별 참가','지역 문제, 비용, 효과 조건 중 무엇을 봤나?'],['9~31분','조별 정책·사건 대응','4턴 무렵 역할 교대. 필요할 때만 전체 일시정지.'],['31~37분','종료·활동지 정리','목표 달성보다 선택과 결과의 설명을 준비.'],['37~46분','서로 다른 2~3개 조 비교','같은 턴에서 차이가 생긴 조건은?'],['46~50분','현실과 모형 구분·개별 정리','게임이 생략한 조건은 무엇인가?']])+
'<h2>비교 조건을 통제하기</h2><p>첫 수업은 같은 시나리오·목표·시드로 통일합니다. 주거·일자리와 생활서비스 등 탐색 방향은 제안하되 정책 순서를 정답으로 주지 않습니다. 같은 시드라도 도시 상태와 선택에 따라 사건이 달라질 수 있습니다.</p>'+box('공정한 평가','목표가 정착 기반이면 서비스 중심 선택은 결말 기준에 불리할 수 있습니다. 이를 학습 실패로 채점하지 않습니다. 과제는 “왜 그 선택을 했고 어떤 조건에서 어떤 결과가 났는가”입니다. 다른 목표·시나리오를 비교할 때는 출발 조건 차이도 명시합니다.')+
'<h2>수업 전 최소 리허설</h2><ul><li>학생 기기1대로 교사 PC 내부망 주소 접속, 참가, 첫 정책, 차트 갱신을 확인.</li><li>교사 서버 창을 유지하고 PC 절전·Wi-Fi 변경을 피함. 기기 격리/방화벽은 학교 환경에서 확인.</li><li>학생에게는 주소와 참가 코드만 공유. 교사용 생성 암호는 공유하지 않음.</li><li>노트북·학생 브라우저에서 글씨와 지도 크기를 확인. 먼저 마친 조는 분석 활동을 확장하고 재도전을 필수로 삼지 않음.</li></ul>'))
teacher.append(section('교사용 02','차트로 결과의 원인을 묻기', '<img class="chart" src="teacher-chart.png" alt="두 조의 돌봄 충족률 추세가 같은 턴에서 비교되는 실제 교사 화면"><p class="note">그림은 자동화 검증용 두 조의 실행 예시이며 실제 학생 결과가 아닙니다. 특정 정책의 보편적 효과나 모범 답안으로 제시하지 않습니다.</p>'+table(['학생의 판단','교사의 후속 질문 / 해석 방향'],[
['“센터를 지었는데 기대만큼 이용하지 않아요.”','정원과 실제 이용을 구분했나? 인력·접근성·수요·운영비를 함께 보게 한다.'],['“빈집을 정비했는데 사람이 많이 오지 않아요.”','주택 외에 채용 가능한 일자리와 정착 조건이 있는가? 단순한 카드 조합 보너스로 설명하지 않는다.'],['“인구는 줄었는데 돌봄 충족률이 올랐어요.”','공급·접근성 개선인가, 수요 감소인가, 둘 다인가? 이용자 수와 수요 분모도 확인한다.'],['“우리 조 예산이 더 많으니 더 성공했어요.”','투자를 적게 했는가, 지속 운영비를 줄였는가? 서비스와 취업 변화도 함께 비교한다.'],['“다른 조보다 더 좋은 정책인데 결과가 달라요.”','같은 턴인가? 대상 생활권의 기존 시설·수요·일자리·주택 조건이 같은가?']])+
'<h2>집계 화면 읽기</h2><p>추세 차트는 0턴의 투자 전 상태와 각 턴 결산입니다. 현재 상태 표는 턴 중 투자비가 먼저 반영될 수 있습니다. 조별 진행 속도가 다르면 현재 수치만으로 비교하지 않습니다.</p><p>전체 돌봄 충족률은 <b>이용 합 ÷ 수요 합</b>, 의료 접근성은 <b>인구 가중 평균</b>입니다. 조별 비율의 단순 평균으로 해석하지 않습니다. 수업 후 CSV·JSON을 내보내고 필요에 따라 수업방을 종료·삭제합니다.</p>'))
teacher.append(section('교사용 03','학습 내용 점검과 자료의 출처', table(['평가 관점','확인할 증거'],[
['문제 발견','지역의 실제 지표와 지도 상태를 근거로 문제를 설명한다.'],['정책 판단','비용·기간·유지비와 효과의 조건을 함께 고려한다.'],['인과 설명','전후 수치를 제시하고 선택→조건 변화→결과로 설명한다.'],['현실 확장','게임의 단순화를 지적하고 권리·형평·세대 간 책임을 논의한다.']])+
'<p>각 항목은 “근거가 구체적임 / 일부 제시 / 추가 설명 필요”로 피드백할 수 있습니다. 게임 승패나 인구 증가량 자체를 점수로 바꾸지 않습니다.</p><h2>사진 내용과 게임의 연결 범위</h2>'+table(['사진에서 확인한 내용','자료·게임에서 다루는 범위'],[
['인구 과잉: 식량·일자리·시설·주택 부족, 자원 부양력','사전 학습에 포함. 게임은 지역 수요·공급을 다루지만 식량 생산·빈곤·국가 인구 억제 정책은 직접 구현하지 않음.'],['저출생·고령화의 원인·영향·정책','사전 학습과 토론에 포함. 게임은 돌봄·주거·고용·의료 접근을 다루지만 실제 출산 결정의 복합성을 재현하지 않음. 돌봄 개선이 출생을 자동으로 늘리는 규칙은 없음.'],['우리나라 지역 격차: 수도권 집중·비수도권 유출','3생활권의 주거·일자리·서비스 차이와 연결. 특정 실제 지역을 그대로 재현한 것은 아님.'],['세대 간 정의·일생활 균형·양성평등·노인 경험과 참여','사진 내용을 바탕으로 토론·평가 질문에 반영. 모두 독립적인 게임 지표나 정책 버튼으로 구현된 것은 아님.']])+
'<h2>출처와 확인 상태</h2><ul><li><b>확인:</b> 사용자가 제공한 사진의 “인구 문제의 해결 방안”, “우리나라 인구 문제의 해결 방안” 표를 읽고 요약·재구성했습니다. 원본 사진은 이 배포 파일에 삽입하지 않았습니다.</li><li><b>확인:</b> 다온시 6단계 엔진과 실시간 수업 모드의 코드·검증 화면. 정책 비용·기간·유지비 표는 엔진 데이터에서 추출했습니다.</li><li><b>미확인:</b> 첨부 파일 「미래엔_통합사회2_교사용교과서_5단원.pdf」는 전송 도구의32MB 한도를 초과하여 원문을 열지 못했습니다. 해당 교과서의 쪽수·문구·도표와 대조했다고 주장하지 않습니다.</li></ul>'+box('교과서 대조를 위한 다음 자료','인구 문제·해결 방안에 해당하는 부분만 별도 PDF로 저장해32MB 미만으로 첨부하거나, 해당 쪽 사진을 제공하면 쪽수와 표현을 대조해 수정할 수 있습니다. 현재는 사진 기반 검토본입니다.')))
css='''@page{size:A4;margin:0}*{box-sizing:border-box}body{margin:0;background:#e4e9df;color:#243f36;font-family:"Noto Sans CJK KR",sans-serif;font-size:10pt;line-height:1.58}.page{width:210mm;height:297mm;padding:15mm 16mm 18mm;background:#fffef8;position:relative;break-after:page;overflow:hidden}.page:last-child{break-after:auto}.eyebrow{font-size:9pt;color:#697d69;letter-spacing:1px;border-top:3px solid #345b47;padding-top:7px}h1{font-size:23pt;letter-spacing:-1px;line-height:1.28;margin:10px 0 17px}h2{font-size:12pt;margin:15px 0 6px;border-left:4px solid #b59a52;padding-left:8px}p{margin:6px 0}table{width:100%;border-collapse:collapse;font-size:9pt;line-height:1.5;margin:10px 0;table-layout:fixed}th,td{border:1px solid #c9d1bd;padding:7px;vertical-align:top;overflow-wrap:anywhere}th{background:#e6ecd9;text-align:left;color:#254d3c}td:first-child{font-weight:600}table th:first-child{width:22%}table:has(th:nth-child(5)) th:first-child{width:24%}table:has(th:nth-child(5)) th:nth-child(2),table:has(th:nth-child(5)) th:nth-child(3),table:has(th:nth-child(5)) th:nth-child(4){width:12%}aside,.goals,.exit{background:#edf0e3;padding:10px 13px;border-left:4px solid #456f58;margin:12px 0}aside p,.goals p{margin:4px 0}.formula{padding:13px;background:#2c5142;color:#fff9e9;text-align:center;font-size:14pt;margin:12px 0}.formula small{font-size:10pt}.lead{font-size:12pt}.note,small{font-size:8.5pt;color:#657163}ul,ol{padding-left:22px;margin:8px 0}li{margin:5px 0}.game{display:block;width:100%;height:70mm;object-fit:contain;margin:8px auto}.chart{display:block;width:100%;height:86mm;object-fit:contain}.name{border-bottom:1px solid #aebda5;padding:8px 0;margin:10px 0}.lines p{height:9mm;border-bottom:1px solid #c1c9b6;margin:0}.footer{position:absolute;bottom:8mm;left:16mm;right:16mm;display:flex;justify-content:space-between;font-size:7.5pt;color:#7a8172;border-top:1px solid #d5daca;padding-top:5px}@media screen{.page{margin:20px auto;box-shadow:0 4px 20px #0002}}'''
def write(name,pages):
 content='<!doctype html><html lang="ko"><meta charset="utf-8"><title>'+name+'</title><style>'+css+'</style><body>'
 for i,p in enumerate(pages,1):content+=f'<section class="page"><div class="eyebrow">다온시 · 통합사회2 · {p["k"]}</div><h1>{p["title"]}</h1>{p["body"]}<footer class="footer"><span>사진 기반 검토본 · 교과서 원문 대조 전 · 2026.10.20</span><span>{i} / {len(pages)}</span></footer></section>'
 (root/(name+'.html')).write_text(content+'</body></html>')
 doc=Document();sec=doc.sections[0];sec.top_margin=Cm(1.5);sec.bottom_margin=Cm(1.5);sec.left_margin=sec.right_margin=Cm(1.6)
 for sty in ['Normal','Heading 1','Heading 2']:
  style=doc.styles[sty];style.font.name='맑은 고딕';style.font.size=Pt(10 if sty=='Normal' else 18 if sty=='Heading 1' else 12);style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
 for i,p in enumerate(pages):
  if i:doc.add_page_break()
  doc.add_paragraph(p['k']);doc.add_heading(p['title'],0)
  soup=BeautifulSoup(p['body'],'html.parser')
  for e in soup.contents:
   if not getattr(e,'name',None):continue
   if e.name=='table':
    trs=e.select('tr');t=doc.add_table(rows=0,cols=len(trs[0].find_all(['th','td'],recursive=False)));t.style='Table Grid'
    for tr in trs:
     cells=t.add_row().cells
     for cell,td in zip(cells,tr.find_all(['th','td'],recursive=False)):cell.text=td.get_text(' ',strip=True)
   elif e.name=='img':doc.add_picture(str(root/e['src']),width=Cm(15))
   elif e.name=='h2':doc.add_heading(e.get_text(),2)
   elif e.name in ['ul','ol']:
    for li in e.find_all('li',recursive=False):doc.add_paragraph(li.get_text(' ',strip=True),style='List Bullet' if e.name=='ul' else 'List Number')
   elif 'lines' in e.get('class',[]):
    for _ in e.find_all('p'):doc.add_paragraph('________________________________________________________________')
   else:doc.add_paragraph(e.get_text(' ',strip=True))
 foot=sec.footer.paragraphs[0];foot.text='사진 기반 검토본 · 교과서 원문 대조 전 · 2026.10.20'
 doc.save(root/(name+'.docx'))
write('학생용_사전학습_안내_활동지',student)
write('교사용_수업운영_해설',teacher)
(root/'pages.json').write_text(json.dumps({'student':student,'teacher':teacher},ensure_ascii=False,indent=2))
print('Created student6 / teacher3 HTML and editableDOCX')
