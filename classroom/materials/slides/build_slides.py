from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION,XL_LABEL_POSITION
import json
ROOT=Path(__file__).resolve().parent
REPO=ROOT.parents[2]
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
C={'bg':'F6F4E9','ink':'213C33','green':'345F49','muted':'637666','line':'CCD4C1','gold':'C5A45A','pale':'E5EAD8','red':'A65A43','white':'FFFFFF','blue':'3A668B'}
FONT='Noto Sans CJK KR'
manifest=[]
def rgb(s):return RGBColor.from_string(C.get(s,s))
def box(s,x,y,w,h,color='pale',line=None,shape=MSO_SHAPE.RECTANGLE):
 a=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h));a.fill.solid();a.fill.fore_color.rgb=rgb(color);a.line.fill.background() if line is None else None
 if line:a.line.color.rgb=rgb(line)
 return a

def text(s,x,y,w,h,value,size=24,color='ink',bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=Inches(.03);tf.margin_right=Inches(.03);tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(value.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(7)
 return sh

def slide(section,title,notes):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('bg');box(s,0,0,.16,7.5,'green');text(s,.55,.3,11,.3,'DAON CITY  /  통합사회2  /  '+section,12,'muted',True);text(s,.55,.88,12.2,.95,title,32,bold=True);box(s,.58,7.05,12.15,.012,'line');text(s,.58,7.16,11.5,.25,'인구 문제와 해결 방안 · 교육용 도시 운영 모형',10,'muted');text(s,12.12,7.12,.6,.3,str(len(prs.slides)),12,'muted');s.notes_slide.notes_text_frame.text=notes;manifest.append({'slide':len(prs.slides),'section':section,'title':title,'notes':notes});return s

def card(s,x,y,w,h,tag,title,body,color='pale',title_size=25,body_size=21):
 box(s,x,y,w,h,color);text(s,x+.2,y+.17,w-.4,.35,tag,13,'muted',True);text(s,x+.2,y+.67,w-.4,1.05,title,title_size,bold=True);text(s,x+.2,y+1.85,w-.4,h-1.95,body,body_size)

def band(s,message,y=6.16,color='green',size=23):
 box(s,.6,y,12.1,.65,color);text(s,.8,y+.1,11.7,.5,message,size,'white',True)

s=slide('문제 제시','인구가 줄어도,\n살기 좋은 도시가 될 수 있을까?', '도입 1분. 질문만 먼저 제시하고 찬성·반대 손들기 또는 짝 대화로 생각을 드러낸다. 정답을 확정하지 않는다. 수업 후 같은 질문으로 돌아온다. 총 17장 중 16~17장은 참고용이다.');
# Replace large title height for two lines.
s.shapes[2].height=Inches(1.5)
text(s,.7,3.0,7.7,1.35,'같은 도시에서 출발합니다.\n어디에 무엇을 먼저 투자할까요?',27)
box(s,9,2.8,3.45,2.45,'green');text(s,9.25,3.05,3,.75,'8년의 선택',32,'white',True);text(s,9.25,4.05,3,1,'주거 · 일자리\n돌봄 · 생활서비스',22,'white')
band(s,'오늘의 과제: 내가 만든 도시의 변화와 그 이유를 설명하기')

s=slide('문제 제시','여러분이라면 어느 도시를 선택하겠습니까?', '도입 2분. 아래는 토의용 가상 사례로 실제 통계나 게임 결과가 아니다. A는 인구·취업 증가에 강점, B는 돌봄에 강점이 있다. 어느 하나를 정답으로 정하지 않는다. 학생이 선택한 기준과 남은 문제를 질문한다.');
card(s,.6,2,5.88,3.45,'도시 A','사람과 일자리가 늘어나는 도시','인구  10,000 → 10,200명\n취업  +120명\n돌봄 수요 충족  65%',body_size=23)
card(s,6.7,2,6,3.45,'도시 B','돌봄이 가까워지는 도시','인구  10,000 → 9,700명\n취업  −30명\n돌봄 수요 충족  90%',body_size=23)
text(s,.7,5.57,12,.35,'토의용 가상 사례입니다. 실제 통계·게임 결과가 아닙니다.',14,'muted')
band(s,'무엇을 기준으로 골랐나요? 아직 알 수 없는 정보는 무엇인가요?')

s=slide('학습 목표','“잘 운영했다”는 것을 근거로 설명하기', '목표 안내 30초. 게임 승패나 인구 증가량을 점수로 평가하지 않는다. 근거 있는 문제 발견, 선택의 효과와 한계, 지속가능성을 관찰한다.');
for i,(a,b,c) in enumerate([('01','문제 발견','지역의 지표와 지도를 보고\n무엇이 부족한지 찾는다.'),('02','선택과 결과','정책 전후 수치를 비교해\n효과와 제한 조건을 설명한다.'),('03','지속가능성','인구 증가와 생활 여건을\n함께 고려해 대안을 제안한다.')]):card(s,.6+i*4.1,2.05,3.9,3.85,a,b,c)
band(s,'핵심 문장: “우리는 ○○를 선택했고, △△ 조건이 바뀌었습니다.”',size=22)

s=slide('학습 정보','인구 변화에는 두 가지 흐름이 있습니다', '개념 확인 1분. 총인구는 규모, 인구 증감은 기간 중 변화다. 예시는 계산 연습용 가상 수치이며 실제 한국 인구 통계나 게임 출생률이 아니다. 여유가 있으면 40명 증가의 이유를 설명하게 한다.');
text(s,.7,1.95,12,.7,'인구 증감 = 자연적 증감 + 사회적 증감',31,bold=True)
card(s,.6,2.95,5.9,2.6,'자연적 증감','출생 − 사망','예: 70 − 110 = −40명',body_size=24)
card(s,6.75,2.95,5.95,2.6,'사회적 증감','전입 − 전출','예: 180 − 100 = +80명',body_size=24)
text(s,.7,5.6,12,.3,'계산용 가상 예시 · 총인구는 전년보다 40명 증가',15,'muted')
band(s,'출생이 사망보다 적어도, 전입이 더 많으면 인구는 늘 수 있습니다.',size=21)

s=slide('학습 정보','인구 문제는 지역마다 다른 모습으로 나타납니다', '사진에 제시된 인구 과잉, 저출생·고령화, 지역 격차를 재구성했다. 공급 부족을 인구수 자체의 문제로만 보지 않는다. 국가별 차이가 있으므로 개발도상국/선진국을 고정적으로 구분하지 않는다. 게임의 세 시나리오는 이런 조건을 단순화한 가상 도시다.');
for i,(a,b,c) in enumerate([('수요가 빠르게 증가','주거·시설 부족','주택·일자리·기반 시설의\n공급이 수요를 따라가는가?'),('사람이 다른 곳으로 이동','공실·지역 격차','지역의 일자리·교육·서비스\n조건이 이동과 어떻게 연결되는가?'),('연령 구성이 변화','돌봄·의료 수요 변화','누가, 어디에서, 어떤\n서비스를 필요로 하는가?')]):card(s,.6+i*4.1,2.05,3.9,3.95,a,b,c,title_size=24,body_size=20)
band(s,'인구수뿐 아니라 “주민의 필요와 공급 조건”을 함께 살펴봅시다.',size=21)

s=slide('학습 정보 · 보충','저출생·고령화는 한 가지 원인으로 설명되지 않습니다', '보충 슬라이드. 사전 학습을 마쳤다면 생략 가능. 저출생을 여성의 사회 참여 탓으로 단정하지 않는다. 고령화는 낮은 출생 수준과 수명 연장 등이 관련된다. 고령 주민은 부담만이 아닌 권리와 경험의 주체다. 첨부 사진의 2018년 고령사회 진입은 과거 서술이며 현재 단계의 최신 통계를 확인한 자료가 아니다.');
card(s,.6,2.0,5.9,3.85,'저출생을 살펴볼 조건','일·돌봄·생활의 조건','주거·양육·교육 부담\n고용 여건과 일·생활 균형\n돌봄 분담, 가족·삶의 가치관',body_size=22)
card(s,6.75,2,5.95,3.85,'고령화에 대응할 조건','일상과 참여의 조건','의료·돌봄·이동 접근성\n노후 소득과 희망하는 일자리\n경험을 살리는 지역사회 참여',body_size=22)
band(s,'특정 성별이나 연령 집단을 문제의 원인으로 단정하지 않습니다.',size=21)

s=slide('학습 정보','정책의 효과는 “필요한 조건”이 있어야 나타납니다', '핵심 인과 슬라이드 1분. 카드 조합 보너스가 아니라 부족한 조건을 해결하는 과정이다. 주택 공급과 입주, 일자리 수와 취업, 시설 정원과 실제 이용을 구분한다. 정책의 효과가 무조건 즉시 인구 증가로 이어진다고 설명하지 않는다.');
for y,left,right in [(2.0,'주택을 정비했다','일자리와 정착 여건이 있어야 입주로 연결'),(3.27,'일자리를 마련했다','노동 공급과 주거 여력이 있어야 정착에 도움'),(4.54,'돌봄센터를 완공했다','운영 인력·접근성·수요가 실제 이용을 제한')]:
 box(s,.6,y,3.6,.95,'pale');text(s,.8,y+.2,3.2,.6,left,24,bold=True);text(s,4.38,y+.14,.7,.6,'→',30,'green',True);text(s,5.1,y+.12,7.25,.8,right,24)
band(s,'선택의 핵심: 지금 가장 부족한 조건은 무엇인가요?',size=24)

s=slide('판단 기준 · 보충','누구의 삶이 좋아지고, 어떤 부담이 남을까요?', '제공 사진의 세대 간 정의, 양성평등 의식, 가족 친화적 직장 문화, 노인의 경험과 지혜 존중을 토론 기준으로 재구성했다. 게임에 모두 별도 수치로 구현된 것은 아님을 설명한다.');
items=[('현재와 미래','지금의 혜택과 이후 유지비를 함께 보기'),('돌봄과 평등','돌봄 책임을 특정 성별에만 맡기지 않기'),('접근성과 형평','시설이 있어도 이용하기 어려운 주민 살피기'),('참여와 존중','고령 주민·이주민의 권리와 참여 고려하기')]
for i,(a,b) in enumerate(items):
 x=.6+(i%2)*6.15;y=2.0+(i//2)*1.8;box(s,x,y,5.9,1.5,'pale');text(s,x+.18,y+.15,5.5,.5,a,25,bold=True);text(s,x+.18,y+.78,5.5,.6,b,20)
band(s,'성장은 인구 증가만으로, 성공은 돈의 크기만으로 판단하지 않습니다.',size=20)

s=slide('게임 안내','같은 도시, 같은 출발 조건, 서로 다른 선택', '게임 진입 안내. 교사가 동일 시나리오·목표·시드로 방을 만든다. 한 조에 조작 기기 하나를 권장한다. 동일 시드라도 정책과 도시 상태에 따라 사건은 달라질 수 있다. 단일 도시 공동 편집 모드가 아니다.');
card(s,.6,2.0,3.9,3.8,'시간','1턴 = 1년','8턴 동안 운영하고\n결과를 비교합니다.',body_size=23)
card(s,4.7,2.0,3.9,3.8,'우리 조','2인 1조','운영자: 조작·정책 제안\n분석가: 비용·결과 확인',body_size=21)
card(s,8.8,2.0,3.9,3.8,'공통 조건','도시·목표·시드','4턴 뒤 역할 교대\n중요한 선택 2회 기록',body_size=22)
band(s,'게임 목표 달성보다 “선택과 결과를 설명하는 것”이 오늘의 과제입니다.',size=20)

s=slide('게임 안내','도시 화면에서 먼저 볼 곳', '교사 실제 게임 시연과 함께 사용한다. 화면은 개발 검증용 예시이며 학생 실제 결과가 아니다. 원본 게임 화면 글씨를 PPT에서 모두 읽기보다 우측 세 가지 위치를 짚어 주고 실제 브라우저에서 시연한다.');
s.shapes.add_picture(str(ROOT.parent/'student-city.png'),Inches(.6),Inches(1.95),width=Inches(8.2),height=Inches(5.1))
for y,title,body in [(2.05,'① 상단 지표','현재 예산 · 인구 · 수지'),(3.45,'② 생활권·시설','지역 문제 · 공급과 이용'),(4.85,'③ 정책·다음 해','비용 확인 → 시행 → 결산')]:
 text(s,9.03,y,3.7,.6,title,24,bold=True);text(s,9.03,y+.65,3.7,.85,body,20)

s=slide('게임 안내','정책을 시행하기 전, 네 가지를 확인하세요', '정책 검토 창을 실제로 열어 짧게 시연한다. 투자비는 시행 시, 유지비는 운영 후 결산에 반영된다. 공급량과 인구 전망의 성격이 다르다는 점을 강조한다. 같은 대상의 중복 투자와 연간 사업 제한도 안내한다.');
for i,(a,b,c) in enumerate([('01','투자비','지금 가진 돈으로 시행 가능한가?'),('02','유지비','앞으로 매년 감당할 수 있는가?'),('03','완료 시점','언제부터 효과를 확인할 수 있는가?'),('04','효과의 조건','인력·수요·접근성·일자리는 충분한가?')]):
 x=.6+(i%2)*6.15;y=1.95+(i//2)*1.8;box(s,x,y,5.9,1.55,'pale');text(s,x+.16,y+.12,5.5,.55,a+'  '+b,25,bold=True);text(s,x+.16,y+.85,5.5,.6,c,20)
band(s,'신규 사업은 한 해 1개 · 같은 정책·같은 대상 중복 투자 불가',size=23)

s=slide('활동 시작','우리 조의 첫 번째 결정을 해봅시다', '교사 PC 서버를 실행한 뒤 실제 학생 접속 주소를 편집한다. 참가 코드는 방 생성 후 입력하거나 교사 화면을 함께 띄운다. 예전 GitHub Pages 개인용 주소로 접속하면 집계되지 않는다. 교사용 생성 암호는 화면에 쓰거나 학생에게 공유하지 않는다. 학생은 수업방 시작 전 기다린다. 활동 시간 약23분, 미완료 조는 현재 턴으로 근거를 기록하도록 한다.');
box(s,.6,1.9,12.1,1.35,'pale');text(s,.85,2.02,11.5,.5,'학생 접속 주소: [교사 PC의 접속 주소 입력]',24,bold=True);text(s,.85,2.63,11.5,.5,'참가 코드: [수업방 생성 후 입력]    조 이름: 예) 1조',22)
text(s,.8,3.65,11.8,1.8,'1  생활권의 문제를 지표와 지도에서 찾기\n2  “이 정책이 왜 필요한가?”에 둘이 합의하기\n3  시행 후 같은 장소와 수치가 어떻게 달라졌는지 확인하기',26)
band(s,'선택 기록: 정책·대상·비용·예상 조건·실제 결과를 짧게 적으세요.',size=21)

s=slide('중간 점검','예상한 만큼 달라지지 않았다면?', '4턴 무렵 필요할 때만 전체 일시정지하고 사용한다. 모든 조에 새로운 정책을 강요하지 않는다. 비교 시 서비스 충족률의 분모인 수요도 확인하게 한다. 역할 교대를 안내하고 충분히 관찰할 시간을 준다.');
text(s,.8,2.05,11.7,2.8,'건물은 완공됐는데, 실제 이용도 늘었나요?\n주택은 늘었는데, 채용 가능한 일자리도 있나요?\n서비스 비율이 올랐다면, 공급과 수요 중 무엇이 변했나요?',29)
box(s,.7,5.2,11.9,.75,'pale');text(s,.92,5.35,11.5,.55,'공급 · 인력 · 접근성 · 수요 · 주거 · 일자리 · 재정',25,'green',True)
band(s,'지금 필요한 것은 새 건물일까요, 기존 시설의 운영 개선일까요?',size=22)

s=slide('결과 비교','같은 턴의 추세를 비교하세요', '게임 후 실제 교사 대시보드로 대체하는 것을 권장한다. 이 차트는 자동화 검증에서 일자리 또는 돌봄 사업1개를 선택한 두 경로의 실제 결과다. 각 턴 사건에는 무료 대응을 했으며 모두 동일 기본 시나리오·목표·시드2026. 두 경로 모두 목표를 달성한 모범 전략이 아니다. 정책 한 가지의 일반적인 인과 효과나 실제 지역 통계로 해석하지 않는다. 공급·수요·취업·재정도 함께 확인한다.');
data=json.loads((ROOT.parents[1]/'artifacts/classroom-export.json').read_text());chartdata=CategoryChartData();chartdata.categories=[str(i)+'턴' for i in range(9)]
for t,label in zip(data['teams'],['일자리 우선 · 검증 예시','돌봄 우선 · 검증 예시']):chartdata.add_series(label,[round(h['coverage']*100,1) for h in t['history']])
ch=s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS,Inches(.65),Inches(2.0),Inches(8.3),Inches(3.95),chartdata).chart
ch.has_legend=True;ch.legend.position=XL_LEGEND_POSITION.BOTTOM;ch.legend.include_in_layout=False;ch.legend.font.name=FONT;ch.legend.font.size=Pt(14)
ch.value_axis.minimum_scale=0;ch.value_axis.maximum_scale=100;ch.value_axis.major_unit=25;ch.value_axis.tick_labels.font.size=Pt(14);ch.category_axis.tick_labels.font.size=Pt(14);ch.value_axis.tick_labels.font.name=ch.category_axis.tick_labels.font.name=FONT
for series,col in zip(ch.series,['green','red']):series.format.line.color.rgb=rgb(col);series.format.line.width=Pt(3)
text(s,9.25,2.15,3.35,.65,'돌봄 수요 충족률 (%)',21,bold=True);text(s,9.25,3.05,3.35,2.5,'비교할 때\n• 같은 턴인가?\n• 초기 조건이 같은가?\n• 다른 지표는 어떤가?',23)
text(s,.8,6,11.7,.4,'검증용 두 경로의 실제 게임 기록입니다. 학생 결과·실제 통계·정답 전략이 아닙니다.',14,'muted')
band(s,'비율이 올랐다면 실제 이용량과 수요도 함께 확인하세요.',y=6.42,size=22)

s=slide('학습 정리','우리가 만든 도시를 설명해 봅시다', '마무리 5분. 발표는 주장-근거-이유-한계 순서로 받는다. 게임 목표 미달 조도 근거를 갖춘 설명이면 학습 성취로 인정한다. 마지막 질문은 게임에서 생략한 원인과 권리를 현실에 연결하기 위한 것이며 별도 퀴즈가 아니다.');
for y,k,v in [(1.95,'주장','우리 도시는 ______ 면에서 개선되었습니다.'),(2.94,'근거','______턴의 ______ 지표가 ______에서 ______로 변했습니다.'),(3.93,'이유','______ 선택이 ______ 조건을 바꾸었기 때문입니다.'),(4.92,'한계','하지만 ______ 문제는 남았고, 현실에서는 ______도 필요합니다.')]:
 box(s,.65,y,1.3,.73,'green');text(s,.8,y+.12,1,.5,k,23,'white',True);text(s,2.15,y+.05,10.3,.8,v,23)
band(s,'다시 질문합니다. 인구가 줄어도 살기 좋은 도시가 될 수 있을까요?',size=21)

s=slide('참고 01 · 필요할 때만','정책 비교표 — 숫자는 게임의 가정입니다', '학생이 정책을 고를 때 참고용으로만 사용한다. 6개 정책의 투자비/준비 기간/연간 유지비를 엔진에서 직접 읽어 구성했다. 의료·이동과 정착 연결은 도시 전체에, 나머지는 대상 생활권에 적용된다. 실효 공급과 실제 이용은 다르다. 유지비 -2는 연간2억 절감이라는 뜻이다.');
policies=json.loads((ROOT.parent/'game-reference.json').read_text())['policies'];rows=[['정책','대상','초기 투자','준비 기간','연간 유지비']]+[[p['name'],'생활권' if p['scope']=='district' else '도시 전체',str(p['cost'])+'억','즉시' if p['delay']==0 else str(p['delay'])+'년','2억 절감' if p['upkeep']<0 else str(p['upkeep'])+'억'] for p in policies]
t=s.shapes.add_table(7,5,Inches(.65),Inches(2),Inches(12),Inches(3.75)).table
for c,w in zip(t.columns,[4.9,1.65,1.65,1.7,2.1]):c.width=Inches(w)
for i,row in enumerate(rows):
 for j,val in enumerate(row):
  c=t.cell(i,j);c.text=val;c.fill.solid();c.fill.fore_color.rgb=rgb('green' if i==0 else ('pale' if i%2 else 'bg'));c.margin_left=Inches(.12);c.margin_right=Inches(.08);c.vertical_anchor=MSO_ANCHOR.MIDDLE
  for p in c.text_frame.paragraphs:p.font.name=FONT;p.font.size=Pt(19);p.font.bold=i==0;p.font.color.rgb=rgb('white' if i==0 else 'ink')
text(s,.8,6,11.8,.9,'건설은 동시에 1개 · 중지 가능한 운영 정책도 투자비는 반환되지 않음\n효과는 수요·인력·접근성·주거·일자리 조건에 따라 달라집니다.',20)

s=slide('참고 02 · 교사용','자료의 근거와 모형의 한계', '수업 전 교사가 확인하는 참고 슬라이드. 학생에게 길게 읽힐 필요는 없다. 사용자 제공 사진 내용을 요약·재구성했으며 원본 사진과 교사용 교과서 PDF를 배포하지 않는다. 교과서 PDF는32MB 전송 한도로 열지 못했으므로 쪽수/도표를 확인한 것처럼 기재하지 않았다. 자료 개정 시 관련 쪽을 확인해야 한다.');
text(s,.75,1.95,11.8,4.6,'확인한 자료\n• 제공된 사진의 인구 문제·해결 방안, 실제 게임 규칙과 검증 기록\n• 정책표는 게임 데이터, 비교 차트는 검증용 두 조의 실행 기록\n\n모형의 한계\n• 실제 출산 결정의 복합성, 식량·국가 인구 정책 등을 모두 구현하지 않음\n• 정책 수치는 실제 정책 효과의 예측값이 아님\n• 고령 주민·이주민을 문제나 자원으로만 보지 않기',22)
text(s,.8,6.35,11.7,.55,'첨부 교과서 PDF는 용량 제한으로 원문 미대조 · 해당 쪽 확인 후 개정 필요',16,'muted')
prs.core_properties.title='같은 출발, 서로 다른 도시 — 통합사회2 공개수업'
prs.core_properties.subject='인구 문제와 해결 방안: 문제 제시·학습 정보·도시 운영·결과 비교'
prs.core_properties.author='다온시 수업 자료'
prs.core_properties.comments='사진 내용과 게임 규칙 기반. 첨부 교과서 PDF 원문 미대조.'
# Validate positions. Some text line boxes intentionally have internal empty space.
for i,s in enumerate(prs.slides,1):
 for sh in s.shapes:
  assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=prs.slide_width+1000 and sh.top+sh.height<=prs.slide_height+1000,(i,sh.shape_id)
prs.save(ROOT/'daon-lesson.pptx')
(ROOT/'slide-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
(ROOT/'발표자_진행노트.txt').write_text('\n\n'.join(f'{m["slide"]}. {m["title"]}\n{m["notes"]}' for m in manifest),encoding='utf-8-sig')
print('Created',len(prs.slides),'slides with presenter notes, native text and editable chart')
