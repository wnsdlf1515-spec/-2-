from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import html,json
root=Path(__file__).parent
meta=[('일    시','____. __. __. __교시'),('장    소','__학년 __반 교실'),('단    원','5. 미래와 지속가능한 삶\n인구 문제와 해결 방안'),('대    상','__학년 __반 학생'),('지도교사','____________')]
goals=['인구 변화와 지역별 생활 여건을 근거로 도시의 주요 문제를 설명할 수 있다.','정책 선택 전후의 지표를 비교하여 정책의 효과와 제한 조건을 설명할 수 있다.','인구 증가와 생활 여건 개선을 구분하여 지속가능한 도시 운영 방안을 제안할 수 있다.']
rows=[['도입\n5분', '문제 제시·\n목표 확인', '• 가상의 두 도시를 비교한다.\n• “인구가 줄어도 살기 좋은 도시가 될 수 있을까?”를 묻는다.\n• 정책의 효과와 한계를 근거로 설명하는 학습 목표를 안내한다.', '• 짝과 도시를 선택하고 이유를 말한다.\n• 판단에 더 필요한 정보를 제안한다.\n• 탐구 질문과 학습 목표를 확인한다.', 'PPT 1~3\n가상 도시 사례', '두 사례의 정답을 확정하지 않는다. 인구 증가만으로 도시를 평가하지 않는다.'], ['전개\n40분', '핵심 개념·\n조작 안내\n7분', '• 자연적·사회적 증감과 지역별 인구 문제를 설명한다(4분).\n• 지역→정책→시간 진행을 시연한다(3분).\n• “보유 예산과 연간 수지는 어떻게 다를까?”를 묻는다.', '• 출생·사망과 전입·전출을 구분한다.\n• 동일한 도시·목표·시드로 접속한다.\n• 2인 1조로 조작자·분석자를 정한다.\n• 비용·기간·유지비·효과 조건을 확인한다.', 'PPT 4·5·7\nPPT 9~12\n조별 기기', '20명·10개 조 기준. 핵심 개념과 조작을 짧게 안내하고 정답 투자 순서는 제시하지 않는다.'], ['', '도시 운영·\n중간 점검\n23분', '• 초반 8분: 지역 진단과 첫 투자를 돕는다. “왜 이 지역인가?”\n• 점검 3분: “기대한 변화와 실제가 같은가?”를 묻고 역할 교대를 안내한다.\n• 후반 12분: 사건 대응과 전략 수정을 돕고 종료 기록을 안내한다.', '• 문제와 근거 수치 1개를 기록하고 첫 투자에 합의한다.\n• 공사·운영·실제 이용을 구분해 확인한다.\n• 예상과 실제를 비교하고 4턴에 역할을 바꾼다.\n• 8턴 운영 후 성과 2개와 남은 문제 1개를 정리한다.', '게임·기록지\nPPT 13\n교사 현황', '1턴=1년. 지연 효과와 운영비를 고려한다. 늦은 조는 동일한 중간 턴으로 비교한다. 종료 후 새 게임을 바로 시작하지 않는다.'], ['', '결과 비교·\n원인 설명\n10분', '• 같은 턴의 서로 다른 투자 경로 2~3개 조를 비교한다.\n• “순서와 위치가 어떤 차이를 만들었나?”를 묻는다.\n• 시행 기록과 사건 대응도 함께 확인한다.', '• 문제→정책·지역→수치 변화→한계 순서로 발표한다.\n• 서비스 비율의 공급·수요를 함께 확인한다.\n• 다른 조와 투자 위치·순서를 비교하고 다음 대안을 제안한다.', '실시간 차트\n정책 기록\nPPT 14', '검증 예시 차트보다 당일 집계를 우선한다. 차트 차이만으로 인과를 단정하지 않는다. 인구 감소 중 생활 개선도 다룬다.'], ['정리·\n차시 예고\n5분', '개념 정리·\n개별 평가', '• 처음 탐구 질문으로 돌아간다.\n• “어떤 근거로 답할 수 있을까?”\n• 주거·일자리·서비스·재정의 연결과 모형의 한계를 정리한다.\n• 실제 지역 자료 탐구를 예고한다.', '• 주장·수치 근거·이유·한계를 개별 작성한다.\n• 다음 도전에서 바꿀 선택을 제안한다.\n• 현실의 일·생활 균형, 돌봄 분담 등 모형 밖 조건을 생각한다.', 'PPT 15\n개별 정리 문항', '게임 승패가 아닌 근거와 해석을 평가한다. 차시 예고는 실제 진도에 맞게 조정한다.']]
notes='수업 운영 기준: 50분, 학생 약 20명, 2인 1조(약 10개 도시). 수업 전 학생 기기 1대로 접속·정책·차트 갱신을 확인한다. 접속 장애 시 교사 시연 화면과 준비된 비교 자료를 활용한 모둠 토의로 전환한다.'
assessment='관찰·활동지 평가: 문제 발견의 근거, 비용과 효과 조건의 고려, 전후 수치에 근거한 인과 설명, 현실의 지속가능성에 대한 제안. 게임 승패 자체는 평가 점수로 사용하지 않는다.'
check='확인 사항: 단원 내 차시 번호는 수업 진도에 맞게 기입. 첨부 교과서 PDF는 전송 용량 제한으로 원문 미대조이며, 소단원명은 제공 사진의 “인구 문제와 해결 방안”을 본시 주제로 사용함. 단원 성취기준 코드는 확인되지 않아 임의로 기재하지 않음.'
def esc(t):return html.escape(t).replace('\n','<br>')
css='''@page{size:A4;margin:0}*{box-sizing:border-box}body{margin:0;font-family:"Noto Sans CJK KR",sans-serif;color:#111;font-size:10pt;line-height:1.5}.page{width:210mm;height:297mm;padding:15mm 13mm;position:relative;break-after:page}.page:last-child{break-after:auto}.cover{border:1px solid black;height:264mm;margin:0 4mm;padding:19mm 13mm}.cover h1{text-align:center;font-size:28pt;font-weight:500;border:3px double #111;padding:13px;margin:0 0 44mm}.meta{width:100%;border-collapse:collapse;font-size:12pt}.meta th{width:24%;font-weight:500}.meta th,.meta td{border:1px solid #111;padding:10px;text-align:center}.meta td{white-space:normal}h2{text-align:center;font-size:18pt;margin:0 0 10px}table{border-collapse:collapse;width:100%;table-layout:fixed}th,td{border:1px solid #333;padding:5px;vertical-align:top}th{background:#f0f0f0;text-align:center}.intro{font-size:9pt;margin-bottom:8px}.intro th{width:13%}.process{font-size:8.3pt;line-height:1.48}.process td{overflow-wrap:break-word}.process .stage{text-align:center;font-weight:600}.note{font-size:8pt;margin:8px 0;line-height:1.5}.footer{position:absolute;bottom:8mm;left:13mm;right:13mm;text-align:center;font-size:8pt}.confirm{font-size:8pt;color:#444;margin-top:15mm}'''
cover='<section class="page"><div class="cover"><h1>(통합사회2) 공개수업<br>지도안</h1><table class="meta">'+''.join(f'<tr><th>{esc(k)}</th><td>{esc(v)}</td></tr>' for k,v in meta)+'</table></div><div class="footer">1</div></section>'
intro='<h2>본시 교수-학습 과정안</h2><table class="intro"><tr><th>대단원</th><td colspan="3">5. 미래와 지속가능한 삶</td></tr><tr><th>소단원</th><td>인구 문제와 해결 방안<br>(도시 운영 모형으로 정책 탐구)</td><th>차시</th><td>적용·탐구 차시<br>단원 내 순서 확인 필요</td></tr><tr><th>학습 목표</th><td colspan="3">'+''.join(f'{i}. {g}<br>' for i,g in enumerate(goals,1))+'</td></tr></table>'
t='<table class="process"><colgroup><col style="width:9%"><col style="width:12%"><col style="width:26%"><col style="width:26%"><col style="width:10%"><col style="width:17%"></colgroup><thead><tr><th rowspan="2">학습<br>단계</th><th rowspan="2">수업내용</th><th colspan="2">교수학습활동</th><th rowspan="2">자료</th><th rowspan="2">유의점</th></tr><tr><th>교수활동</th><th>학습활동</th></tr></thead><tbody>'
for i,row in enumerate(rows):
 t+='<tr>'
 for j,c in enumerate(row):
  if j==0 and i in [2,3]:continue
  t+=f'<td'+(' rowspan="3"' if i==1 and j==0 else '')+(' class="stage"' if j==0 else '')+'>'+esc(c)+'</td>'
 t+='</tr>'
t+='</tbody></table>'
body='<section class="page">'+intro+t+f'<p class="note"><b>운영·대응</b> {notes}</p><p class="note"><b>평가</b> {assessment}</p><p class="note"><b>작성 확인</b> {check}</p><div class="footer">2</div></section>'
(root/'수업공개_지도안_작성본.html').write_text('<!doctype html><html lang="ko"><meta charset="utf-8"><title>통합사회2 공개수업 지도안</title><style>'+css+'</style><body>'+cover+body+'</body></html>')
# Editable DOCX, same template fields and merged header/phase cells.
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.5);sec.bottom_margin=Cm(1.5);sec.left_margin=sec.right_margin=Cm(1.3)
normal=doc.styles['Normal'];normal.font.name='맑은 고딕';normal.font.size=Pt(9);normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕');normal.paragraph_format.space_after=Pt(2)
p=doc.add_paragraph();p.paragraph_format.space_before=Pt(45);p.paragraph_format.space_after=Pt(110);p.alignment=1;r=p.add_run('(통합사회2) 공개수업\n지도안');r.font.size=Pt(28)
tb=doc.add_table(rows=0,cols=2);tb.style='Table Grid';tb.columns[0].width=Cm(3.3);tb.columns[1].width=Cm(13)
for k,v in meta:
 cells=tb.add_row().cells;cells[0].text=k;cells[1].text=v
 for c in cells:
  for p in c.paragraphs:p.alignment=1;p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(8)
doc.add_page_break();p=doc.add_paragraph('본시 교수-학습 과정안');p.alignment=1;p.runs[0].font.size=Pt(17)
tb=doc.add_table(rows=3,cols=4);tb.style='Table Grid';tb.cell(0,0).text='대단원';tb.cell(0,1).merge(tb.cell(0,3)).text='5. 미래와 지속가능한 삶';tb.cell(1,0).text='소단원';tb.cell(1,1).text='인구 문제와 해결 방안\n(도시 운영 모형으로 정책 탐구)';tb.cell(1,2).text='차시';tb.cell(1,3).text='적용·탐구 차시\n단원 내 순서 확인 필요';tb.cell(2,0).text='학습 목표';tb.cell(2,1).merge(tb.cell(2,3)).text='\n'.join(f'{i}. {g}' for i,g in enumerate(goals,1))
doc.add_paragraph('')
tb=doc.add_table(rows=2,cols=6);tb.style='Table Grid';tb.autofit=False
widths=[1.6,2.2,4.8,4.8,1.8,3.2]
for c,w in zip(tb.columns,widths):c.width=Cm(w)
for j,name in enumerate(['학습 단계','수업내용','교수활동','학습활동','자료','유의점']):tb.cell(1,j).text=name
for j in [0,1,4,5]:tb.cell(0,j).merge(tb.cell(1,j)).text=['학습 단계','수업내용','','','자료','유의점'][j]
tb.cell(0,2).merge(tb.cell(0,3)).text='교수학습활동'
for row in rows:
 cells=tb.add_row().cells
 for c,text in zip(cells,row):
  c.text=text
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(0)
   for r in p.runs:r.font.size=Pt(8)
tb.cell(3,0).merge(tb.cell(5,0)).text='전개\n40분'
for row in tb.rows:
 trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
for label,text in [('운영·대응',notes),('평가',assessment),('작성 확인',check)]:
 p=doc.add_paragraph(label+' — '+text)
 for r in p.runs:r.font.size=Pt(8)
doc.save(root/'수업공개_지도안_작성본.docx')
plain='(통합사회2) 공개수업 지도안\n\n'+'\n'.join(k+': '+v for k,v in meta)+'\n\n본시 교수-학습 과정안\n대단원:5. 미래와 지속가능한 삶\n소단원:인구 문제와 해결 방안\n차시:적용·탐구 차시(단원 내 순서 확인 필요)\n\n학습 목표\n'+'\n'.join(goals)+'\n\n'
for row in rows:plain+='\n'.join(k+': '+v for k,v in zip(['학습 단계','수업내용','교수활동','학습활동','자료','유의점'],row))+'\n\n'
plain+='운영·대응: '+notes+'\n평가: '+assessment+'\n작성 확인: '+check
(root/'한글_붙여넣기용_본문.txt').write_text(plain,encoding='utf-8-sig')
(root/'README.txt').write_text('첨부 양식의 항목에 맞춰 작성한 공개 다운로드용 지도안입니다. 일시·장소·대상·지도교사 칸은 직접 기입하세요. 원본 HWP를 직접 수정한 파일은 아닙니다.\nDOCX를 한글에서 열고 표·쪽 나눔을 확인한 뒤 HWP로 다른 이름 저장할 수 있습니다. PDF는 확인·인쇄용이고 TXT는 원본 HWP에 붙여넣기용입니다.\n50분 수업으로 구성했으며 단원 내 차시 번호와 실제 진도만 확인해 주세요. 차시 예고는 후속 탐구 제안이므로 실제 수업 계획에 맞게 조정하세요.\n수업 내용은 유지하고 개인 식별 정보를 빈칸으로 처리했습니다. 교과서 원문 대조는 용량 제한으로 미완료입니다.\n',encoding='utf-8-sig')
print('Created revised lesson plan')
